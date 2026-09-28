"""ATS-friendly single-column DOCX rendering for tailored resumes."""

from __future__ import annotations

from collections.abc import Mapping, Sequence
from io import BytesIO
from pathlib import Path
from typing import Any
from uuid import uuid4

from docx import Document
from docx.document import Document as DocumentObject
from docx.shared import Inches, Pt

from .models import MasterResume, ResumeBullet


def render_resume_docx(
    resume: MasterResume,
    *,
    match_result: Mapping[str, Any] | None = None,
    tailored_bullets: Mapping[str, str] | None = None,
) -> bytes:
    """Render a resume as a simple one-column DOCX and return its bytes."""
    match_result = match_result or {}
    tailored_bullets = tailored_bullets or {}
    top_items = match_result.get("top_items", [])
    bullet_scores: dict[str, int] = {}
    project_scores: dict[str, int] = {}
    for item in top_items if isinstance(top_items, list) else []:
        if not isinstance(item, Mapping):
            continue
        score = int(item.get("relevance_score", 0) or 0)
        if item.get("item_type") == "bullet" and item.get("id"):
            bullet_scores[str(item["id"])] = max(score, bullet_scores.get(str(item["id"]), 0))
        elif item.get("item_type") == "project" and item.get("title"):
            project_scores[str(item["title"])] = max(
                score, project_scores.get(str(item["title"]), 0)
            )

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(3)
    for name in ("Title", "Heading 1", "Heading 2"):
        document.styles[name].font.name = "Arial"

    contact = resume.contact
    name_paragraph = document.add_paragraph()
    name_run = name_paragraph.add_run(contact.name)
    name_run.bold = True
    name_run.font.size = Pt(16)
    details = [
        contact.location,
        contact.email,
        contact.phone,
        contact.linkedin,
        contact.github,
        contact.website,
    ]
    contact_line = " | ".join(value.strip() for value in details if value and value.strip())
    if contact_line:
        document.add_paragraph(contact_line)

    if resume.summary.strip():
        document.add_heading("Professional Summary", level=1)
        document.add_paragraph(resume.summary.strip())

    if resume.skills:
        document.add_heading("Skills", level=1)
        matched = {str(skill).casefold() for skill in match_result.get("matched_skills", []) or []}
        for category, skills in resume.skills.items():
            ordered_skills = sorted(
                skills, key=lambda skill: (skill.casefold() not in matched, skill.casefold())
            )
            document.add_paragraph(f"{category}: {', '.join(ordered_skills)}")

    experiences = list(enumerate(resume.experience))
    experiences.sort(
        key=lambda pair: (
            -max((bullet_scores.get(b.id, 0) for b in pair[1].bullets), default=0),
            pair[0],
        )
    )
    projects = list(enumerate(resume.projects))
    projects.sort(
        key=lambda pair: (
            -max(
                project_scores.get(pair[1].name, 0),
                *(bullet_scores.get(b.id, 0) for b in pair[1].bullets),
            ),
            pair[0],
        )
    )

    def write_experience() -> None:
        document.add_heading("Professional Experience", level=1)
        for _, entry in experiences:
            dates = entry.dates.start + (
                f" - {entry.dates.end}" if entry.dates.end else " - Present"
            )
            document.add_heading(f"{entry.role} | {entry.company}", level=2)
            document.add_paragraph(dates)
            _add_bullets(document, entry.bullets, tailored_bullets)

    def write_projects() -> None:
        document.add_heading("Projects", level=1)
        for _, project in projects:
            document.add_heading(project.name, level=2)
            if project.tech:
                document.add_paragraph(f"Technologies: {', '.join(project.tech)}")
            _add_bullets(document, project.bullets, tailored_bullets)

    experience_relevance = max(
        [
            bullet_scores.get(bullet.id, 0)
            for entry in resume.experience
            for bullet in entry.bullets
        ],
        default=0,
    )
    project_relevance = max(
        [
            *project_scores.values(),
            *(
                bullet_scores.get(bullet.id, 0)
                for project in resume.projects
                for bullet in project.bullets
            ),
        ],
        default=0,
    )
    relevant_sections = [
        (experience_relevance, 0, bool(experiences), write_experience),
        (project_relevance, 1, bool(projects), write_projects),
    ]
    for _, _, present, write_section in sorted(
        relevant_sections, key=lambda section: (-section[0], section[1])
    ):
        if present:
            write_section()

    if resume.education:
        document.add_heading("Education", level=1)
        for item in resume.education:
            line = " | ".join(part for part in (item.degree, item.field, item.institution) if part)
            if item.dates:
                line += f" | {item.dates.start}" + (
                    f" - {item.dates.end}" if item.dates.end else ""
                )
            document.add_paragraph(line)
            for detail in item.details:
                document.add_paragraph(detail, style="List Bullet")

    if resume.certifications:
        document.add_heading("Certifications", level=1)
        for item in resume.certifications:
            line = item.name + (f" | {item.issuer}" if item.issuer else "")
            if item.date:
                line += f" ({item.date})"
            document.add_paragraph(line)

    output = BytesIO()
    document.save(output)
    return output.getvalue()


def save_resume_docx(data: bytes, path: str | Path) -> Path:
    """Atomically save rendered DOCX bytes to a path."""
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    temporary = target.with_name(f".{target.name}.{uuid4().hex}.tmp")
    temporary.write_bytes(data)
    temporary.replace(target)
    return target


def _add_bullets(
    document: DocumentObject, bullets: Sequence[ResumeBullet], tailored: Mapping[str, str]
) -> None:
    for bullet in bullets:
        text = tailored.get(bullet.id, bullet.text)
        if text.strip():
            document.add_paragraph(text.strip(), style="List Bullet")
